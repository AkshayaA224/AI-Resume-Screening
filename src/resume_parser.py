
from pathlib import Path

import pymupdf
from docx import Document


def extract_text_from_pdf(file_path):
    """Extract text from a PDF resume."""
    text_parts = []

    with pymupdf.open(file_path) as pdf:
        for page in pdf:
            text_parts.append(page.get_text())

    text = "\n".join(text_parts).strip()

    if not text:
        raise ValueError(
            "No text found in PDF. "
            "It may be a scanned document requiring OCR."
        )

    return text


def extract_text_from_docx(file_path):
    """Extract text from a DOCX resume."""
    document = Document(file_path)
    parts = []

    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            parts.append(paragraph.text.strip())

    # Also extract text from tables commonly used in resumes.
    for table in document.tables:
        for row in table.rows:
            parts.append(
                " | ".join(cell.text.strip() for cell in row.cells)
            )

    text = "\n".join(parts).strip()

    if not text:
        raise ValueError("No text found in DOCX resume.")

    return text


def extract_resume_text(file_path):
    """Choose the appropriate parser based on file extension."""
    path = Path(file_path)
    extension = path.suffix.lower()

    if extension == ".pdf":
        return extract_text_from_pdf(path)

    if extension == ".docx":
        return extract_text_from_docx(path)

    raise ValueError(
        "Unsupported file format. Upload a PDF or DOCX resume."
    )
